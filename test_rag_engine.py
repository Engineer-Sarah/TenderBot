"""
Offline tests for rag_engine.py - no API key, no network, no ChromaDB needed.
Run:  python -m pytest test_rag_engine.py -v      (or: python test_rag_engine.py)

A fake in-memory Chroma client and a hashed bag-of-words embedder replace the
real services, so these tests verify OUR logic (chunking, page tracking,
idempotent upload, thresholds, verdict parsing, fallbacks, error handling).
"""
import io
import json
import math
import re

import pytest
from docx import Document
from reportlab.pdfgen import canvas

import rag_engine as re_mod
from rag_engine import RAGEngine, RAGError, chunk_text

DIM = 128


# ----------------------------- fakes --------------------------------------- #
def fake_embed(texts, task_type):
    out = []
    for t in texts:
        v = [0.0] * DIM
        for w in re.findall(r"\w+", t.lower()):
            v[hash(w) % DIM] += 1.0
        out.append(v)
    return out


class FakeCollection:
    def __init__(self):
        self.rows = {}

    def count(self):
        return len(self.rows)

    def upsert(self, ids, documents, embeddings, metadatas):
        for i, d, e, m in zip(ids, documents, embeddings, metadatas):
            self.rows[i] = (d, e, m)

    def query(self, query_embeddings, n_results, include):
        q = query_embeddings[0]
        scored = []
        for _id, (d, e, m) in self.rows.items():
            cos = sum(a * b for a, b in zip(q, e))
            scored.append((1 - cos, d, m))
        scored.sort(key=lambda x: x[0])
        scored = scored[:n_results]
        return {
            "documents": [[s[1] for s in scored]],
            "metadatas": [[s[2] for s in scored]],
            "distances": [[s[0] for s in scored]],
        }

    def get(self, where=None, include=None):
        ids, metas = [], []
        for _id, (d, e, m) in self.rows.items():
            if where is None or all(m.get(k) == v for k, v in where.items()):
                ids.append(_id)
                metas.append(m)
        return {"ids": ids, "metadatas": metas}

    def delete(self, where):
        for _id in self.get(where)["ids"]:
            del self.rows[_id]


class FakeChroma:
    def __init__(self):
        self.cols = {}

    def get_or_create_collection(self, name, metadata=None):
        return self.cols.setdefault(name, FakeCollection())

    def delete_collection(self, name):
        del self.cols[name]


def make_engine(generate_fn=None):
    return RAGEngine(
        chroma_client=FakeChroma(),
        embed_fn=fake_embed,
        generate_fn=generate_fn or (lambda prompt, system, json_mode=False: "ANSWER"),
    )


def make_pdf(pages):
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for text in pages:
        c.drawString(50, 800, text)
        c.showPage()
    c.save()
    return buf.getvalue()


def make_docx(paragraphs, table_rows=None):
    d = Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table_rows:
        t = d.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ----------------------------- chunking ------------------------------------ #
def test_chunk_sizes_and_overlap():
    text = " ".join(f"w{i}" for i in range(1200))
    chunks = chunk_text(text, 500, 50)
    assert [len(c.split()) for c in chunks] == [500, 500, 300]
    assert chunks[0].split()[-50:] == chunks[1].split()[:50]  # overlap present


def test_chunk_no_tiny_duplicate_tail():
    text = " ".join(f"w{i}" for i in range(500))
    assert len(chunk_text(text, 500, 50)) == 1
    text = " ".join(f"w{i}" for i in range(501))
    assert len(chunk_text(text, 500, 50)) == 2


def test_chunk_edge_cases():
    assert chunk_text("") == []
    assert chunk_text("   \n ") == []
    with pytest.raises(ValueError):
        chunk_text("a b c", 10, 10)
    with pytest.raises(ValueError):
        chunk_text("a b c", 0, 0)


def test_page_aware_chunking():
    pages = [(1, " ".join(["a"] * 300)), (2, " ".join(["b"] * 300))]
    chunks = re_mod._chunk_pages(pages, 500, 50)
    assert chunks[0]["page_start"] == 1 and chunks[0]["page_end"] == 2
    assert chunks[-1]["page_end"] == 2


def test_collection_name_sanitised():
    n = re_mod._safe_collection_name("Ali & Sons (Pvt) Ltd.")
    assert 3 <= len(n) <= 63 and re.fullmatch(r"[a-zA-Z0-9][a-zA-Z0-9._-]*[a-zA-Z0-9]", n)
    assert len(re_mod._safe_collection_name("x" * 200)) <= 63


# ----------------------------- ingestion ----------------------------------- #
def test_upload_pdf_docx_txt_and_list():
    e = make_engine()
    res = e.upload_docs(
        [
            ("pec.pdf", make_pdf(["PEC license category C5 valid until 2027", "Engineering council"])),
            ("audit.docx", make_docx(["Annual turnover PKR 7 crore"], [["Year", "Turnover"], ["2025", "7 crore"]])),
            ("note.txt", b"NTN number 1234567-8 registered with FBR"),
        ]
    )
    assert [r["status"] for r in res] == ["ok", "ok", "ok"]
    assert res[0]["pages"] == 2
    docs = e.list_documents()
    assert {d["source"] for d in docs} == {"pec.pdf", "audit.docx", "note.txt"}
    assert e.stats()["documents"] == 3


def test_docx_table_text_is_indexed():
    e = make_engine()
    e.upload_docs(("t.docx", make_docx(["x"], [["Project", "Value"], ["Road", "12 crore"]])))
    hit = e.search_docs("Road 12 crore")[0]
    assert "Road | 12 crore" in hit["text"] or "12 crore" in hit["text"]


def test_upload_is_idempotent():
    e = make_engine()
    data = b"Company registered under SECP with NTN number"
    e.upload_docs(("a.txt", data))
    e.upload_docs(("a.txt", data))
    assert e.stats()["chunks"] == 1


def test_bad_files_return_errors_not_exceptions():
    e = make_engine()
    res = e.upload_docs(
        [
            ("empty.txt", b""),
            ("bad.pdf", b"not a pdf"),
            ("img.png", b"123"),
            "/nonexistent/file.pdf",
            ("scanned.pdf", make_pdf([""])),
        ]
    )
    assert all(r["status"] == "error" and r["error"] for r in res)
    assert "scanned" in res[4]["error"].lower()
    assert e.stats()["chunks"] == 0


def test_file_like_upload_streamlit_style():
    e = make_engine()
    f = io.BytesIO(b"PEC license C5")
    f.name = "pec.txt"
    assert e.upload_docs(f)[0]["status"] == "ok"


def test_company_isolation():
    e = make_engine()
    e.upload_docs(("a.txt", b"alpha company secret turnover"), company_id="A")
    assert e.search_docs("turnover", company_id="B") == []


# ----------------------------- querying ------------------------------------ #
def test_query_empty_knowledge_base():
    r = make_engine().query_docs("PEC license?")
    assert r["found"] is False and r["confidence"] == "none"


def test_query_empty_question_rejected():
    with pytest.raises(ValueError):
        make_engine().query_docs("   ")


def test_query_relevant_uses_llm_and_cites_page():
    e = make_engine()
    e.upload_docs(("pec.pdf", make_pdf(["PEC license category C5 valid until 2027"])))
    r = e.query_docs("PEC license category C5")
    assert r["found"] and r["answer"] == "ANSWER"
    assert r["sources"][0]["source"] == "pec.pdf" and r["sources"][0]["pages"] == "1"
    assert r["match_percent"] > 55


def test_query_irrelevant_skips_llm_no_hallucination():
    calls = []
    e = make_engine(lambda p, s, json_mode=False: calls.append(1) or "HALLUCINATION")
    e.upload_docs(("pec.txt", b"PEC license category C5 valid until 2027"))
    r = e.query_docs("quantum physics of black holes")
    assert r["found"] is False and calls == []
    assert "nahi mili" in r["answer"]


def test_llm_failure_degrades_gracefully():
    def boom(prompt, system, json_mode=False):
        raise RAGError("quota")

    e = make_engine(boom)
    e.upload_docs(("pec.txt", b"PEC license category C5 valid until 2027"))
    r = e.query_docs("PEC license category C5")
    assert r["answer"] is None and "LLM unavailable" in r["warning"] and r["sources"]


# --------------------------- check_requirement ------------------------------ #
def test_requirement_missing_when_nothing_relevant():
    e = make_engine()
    e.upload_docs(("a.txt", b"NTN number registered with FBR"))
    r = e.check_requirement("blockchain mining rig hardware")
    assert r["status"] == "missing" and r["evidence"] == []


def test_requirement_llm_verdict_parsed_including_code_fence():
    gen = lambda p, s, json_mode=False: '```json\n{"verdict": "partial", "reason": "C6 not C5"}\n```'
    e = make_engine(gen)
    e.upload_docs(("pec.txt", b"PEC license category C6 valid until 2027"))
    r = e.check_requirement("PEC license category C5 valid until 2027")
    assert r["status"] == "partial" and r["reason"] == "C6 not C5"


def test_requirement_bad_llm_json_falls_back_to_similarity():
    e = make_engine(lambda p, s, json_mode=False: "this is not json")
    e.upload_docs(("pec.txt", b"PEC license category C5 valid until 2027"))
    r = e.check_requirement("PEC license category C5 valid until 2027")
    assert r["status"] in {"met", "partial"} and "similarity" in r["reason"]


def test_requirement_invalid_verdict_value_ignored():
    e = make_engine(lambda p, s, json_mode=False: '{"verdict": "banana", "reason": "x"}')
    e.upload_docs(("pec.txt", b"PEC license category C5 valid until 2027"))
    r = e.check_requirement("PEC license category C5 valid until 2027")
    assert r["status"] in {"met", "partial"}


# ------------------------------ management ---------------------------------- #
def test_delete_document_and_clear():
    e = make_engine()
    res = e.upload_docs([("a.txt", b"alpha doc about turnover"), ("b.txt", b"beta doc about experience")])
    assert e.delete_document(res[0]["doc_id"]) == 1
    assert e.delete_document("nonexistent") == 0
    assert [d["source"] for d in e.list_documents()] == ["b.txt"]
    e.clear_company()
    assert e.stats()["documents"] == 0
    e.clear_company()  # clearing twice must not crash


def test_missing_api_key_raises_clear_error(monkeypatch):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    with pytest.raises(re_mod.RAGConfigError):
        re_mod._api_key()


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-v"]))
